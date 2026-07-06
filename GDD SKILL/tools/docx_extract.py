import sys
import os

def extract_docx(docx_path, output_path):
    """Extract text from a docx file using python-docx, falling back to zipfile XML parsing"""
    try:
        from docx import Document
        doc = Document(docx_path)
        with open(output_path, 'w', encoding='utf-8') as f:
            # Extract paragraphs
            for para in doc.paragraphs:
                style = para.style.name if para.style else ''
                text = para.text.strip()
                if text:
                    f.write(f"[{style}] {text}\n")
            # Extract tables
            for ti, table in enumerate(doc.tables):
                f.write(f"\n=== 表格 {ti+1} ===\n")
                for ri, row in enumerate(table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    f.write(' | '.join(cells) + '\n')
        print(f"OK: python-docx -> {output_path}")
    except ImportError:
        # Fallback: docx is a zip file with XML inside
        import zipfile
        from xml.etree import ElementTree as ET
        with zipfile.ZipFile(docx_path) as z:
            xml_content = z.read('word/document.xml')
        tree = ET.fromstring(xml_content)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with open(output_path, 'w', encoding='utf-8') as f:
            for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                texts = []
                for t in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                    if t.text:
                        texts.append(t.text)
                line = ''.join(texts).strip()
                if line:
                    f.write(line + '\n')
        print(f"OK: zipfile fallback -> {output_path}")

if __name__ == '__main__':
    ref_dir = r"g:\GDD SKILL\策划案参考"
    out_dir = r"g:\GDD SKILL\策划案参考"

    files = [
        '_【Slime Legion】新公会战"军团争霸"外围系统.docx',
        '竞技场高级赛（纷乱的群殴锦标赛）—— 功能点梳理.docx',
    ]

    for f in files:
        in_path = os.path.join(ref_dir, f)
        out_name = f.replace('.docx', '.txt')
        out_path = os.path.join(out_dir, out_name)
        try:
            extract_docx(in_path, out_path)
        except Exception as e:
            print(f"FAIL {f}: {e}")
